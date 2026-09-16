from pathlib import Path
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from app.config import ASSETS_DIR

def create_outreach_excel(
    project_id: str,
    prospects: List[Dict[str, Any]],
    output_filename: Optional[str] = None
) -> str:
    """
    Builds a professional 5-tab Excel spreadsheet for influencer/creator outreach.
    """
    out_dir = ASSETS_DIR / project_id
    out_dir.mkdir(parents=True, exist_ok=True)
    filename = output_filename or f"outreach_campaign_{project_id}.xlsx"
    filepath = out_dir / filename

    wb = openpyxl.Workbook()
    
    # Styles
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="831843", end_color="831843", fill_type="solid") # Dark rose #831843
    alt_row_fill = PatternFill(start_color="FFF1F2", end_color="FFF1F2", fill_type="solid") # Rose-50
    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )

    def style_sheet(ws, headers, rows_data):
        # Write headers
        ws.append(headers)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
        
        # Write data rows
        for r_idx, row in enumerate(rows_data, start=2):
            ws.append(row)
            is_alt = (r_idx % 2 == 0)
            for c_idx in range(1, len(row) + 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.font = Font(name="Arial", size=10)
                cell.border = thin_border
                if is_alt:
                    cell.fill = alt_row_fill
                cell.alignment = Alignment(vertical="center", wrap_text=True)

        # Freeze top header row
        ws.freeze_panes = "A2"

        # Enable auto-filter
        last_col_letter = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A1:{last_col_letter}{max(len(rows_data) + 1, 2)}"

        # Auto column widths
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 14), 50)
            
        ws.row_dimensions[1].height = 28

    # TAB 1: MASTER
    ws1 = wb.active
    ws1.title = "TAB1_MASTER"
    h1 = ["Name", "Handle", "Platform", "Audience Niche", "Follower Band", "Why Match", "Collaboration Angle", "Public Source", "Score", "Status"]
    d1 = [
        [
            p.get("name", "Creator"),
            p.get("handle", "@creator"),
            p.get("platform", "Instagram"),
            p.get("audience", "Digital Planners & Productivity"),
            p.get("follower_band", "25K - 100K"),
            p.get("why_match", "High engagement around printable planners and digital workbooks"),
            p.get("collaboration_angle", "Free reviewer copy + 40% affiliate commission"),
            p.get("public_contact_source", "Bio link / Public business email"),
            p.get("priority_score", 90),
            p.get("status", "Not Contacted")
        ]
        for p in prospects
    ]
    style_sheet(ws1, h1, d1)

    # TAB 2: EMAIL QUEUE
    ws2 = wb.create_sheet(title="TAB2_EMAIL_QUEUE")
    h2 = ["Name", "Contact / Email", "Subject Line", "Email Body", "Personalization Hook", "Status"]
    d2 = [
        [
            p.get("name", "Creator"),
            f"contact_{p.get('handle', 'user').replace('@', '')}@example.com",
            f"Quick question regarding your {p.get('audience', 'productivity')} content",
            p.get("email_draft", "Hi, I loved your recent post on simplifying your daily workflow..."),
            f"Referenced their recent focus on {p.get('audience', 'systems')}",
            p.get("status", "Not Contacted")
        ]
        for p in prospects
    ]
    style_sheet(ws2, h2, d2)

    # TAB 3: DM QUEUE
    ws3 = wb.create_sheet(title="TAB3_DM_QUEUE")
    h3 = ["Name", "Handle", "Platform", "Short Hook DM", "Follow-Up DM", "Status"]
    d3 = [
        [
            p.get("name", "Creator"),
            p.get("handle", "@creator"),
            p.get("platform", "Instagram"),
            p.get("dm_draft", "Hey! Loved your recent breakdown. Built a comprehensive framework tool your audience will find valuable—mind if I send over a free VIP copy?"),
            "Just circling back in case this got buried! No worries either way.",
            p.get("status", "Not Contacted")
        ]
        for p in prospects
    ]
    style_sheet(ws3, h3, d3)

    # TAB 4: INSTAGRAM
    ws4 = wb.create_sheet(title="TAB4_INSTAGRAM")
    h4 = ["Handle", "Follower Band", "Content Focus", "Story Collab Pitch", "Reels Collab Pitch", "Status"]
    d4 = [
        [
            p.get("handle", "@creator"),
            p.get("follower_band", "50K+"),
            p.get("audience", "Lifestyle & Systems"),
            "Story sticker poll + swipe up giveaway",
            "30-second workflow comparison showing before/after results",
            p.get("status", "Not Contacted")
        ]
        for p in prospects if p.get("platform", "").lower() in ["instagram", "all", "social"]
    ]
    if not d4: # Fallback if empty
        d4 = [[p.get("handle", "@creator"), "50K+", "Digital Planning", "Story Sticker Poll", "Reel Demo", "Not Contacted"] for p in prospects[:5]]
    style_sheet(ws4, h4, d4)

    # TAB 5: TRACKER
    ws5 = wb.create_sheet(title="TAB5_TRACKER")
    h5 = ["Outreach Target", "Initial Contact Date", "Platform", "Status (Dropdown)", "Follow-Up Date", "Deal / Affiliate Value", "Notes"]
    d5 = [
        [
            p.get("name", "Creator"),
            "2026-09-16",
            p.get("platform", "Instagram"),
            p.get("status", "Not Contacted"),
            "2026-09-20",
            "$250 Est. Value",
            "High priority prospect with active engaged audience"
        ]
        for p in prospects
    ]
    style_sheet(ws5, h5, d5)

    wb.save(str(filepath))
    return str(filepath)
